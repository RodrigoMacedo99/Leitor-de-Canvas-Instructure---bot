from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from pptx import Presentation

from cbot.modules.conteudo.infrastructure.file_storage import FileStorage
from cbot.modules.conteudo.infrastructure.text_extraction import (
    ExtratorDeArquivos,
    FormatoNaoSuportado,
    Registry,
)
from cbot.shared.errors import ValidationFailed
from tests.helpers import make_pdf


def test_storage_hashes_and_isolates_by_student(tmp_path: Path) -> None:
    storage, a, b = FileStorage(tmp_path), uuid4(), uuid4()
    gravado = storage.gravar(a, iter([b"ab", b"cd"]), 100)
    assert gravado.tamanho == 4 and gravado.caminho.read_bytes() == b"abcd"
    assert storage.existe(a, gravado.sha256) and not storage.existe(b, gravado.sha256)


def test_storage_aborts_over_limit_without_residue(tmp_path: Path) -> None:
    storage, aluno = FileStorage(tmp_path), uuid4()
    with pytest.raises(ValidationFailed):
        storage.gravar(aluno, iter([b"x" * 10]), 5)
    assert list((tmp_path / str(aluno)).rglob("*.part")) == []
    assert [p for p in (tmp_path / str(aluno)).rglob("*") if p.is_file()] == []


def test_storage_deletes_all_of_a_student(tmp_path: Path) -> None:
    storage, aluno = FileStorage(tmp_path), uuid4()
    storage.gravar(aluno, iter([b"z"]), 10)
    storage.apagar_aluno(aluno)
    assert not (tmp_path / str(aluno)).exists()


def test_pdf_extraction_keeps_page_numbers(tmp_path: Path) -> None:
    f = tmp_path / "a.pdf"
    f.write_bytes(make_pdf("Camada de transporte"))
    secoes = ExtratorDeArquivos().extrair(f, "application/pdf", "a.pdf")
    assert secoes[0].pagina == 1 and "Camada de transporte" in secoes[0].texto


def test_docx_splits_on_headings(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("TCP", level=1)
    doc.add_paragraph("Orientado a conexão.")
    doc.add_heading("UDP", level=1)
    doc.add_paragraph("Sem conexão.")
    f = tmp_path / "a.docx"
    doc.save(f)
    secoes = ExtratorDeArquivos().extrair(f, "", "a.docx")
    assert [s.titulo for s in secoes] == ["TCP", "UDP"]
    assert "Sem conexão" in secoes[1].texto


def test_pptx_one_section_per_slide(tmp_path: Path) -> None:
    prs = Presentation()
    for titulo in ("Slide um", "Slide dois"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = titulo
    f = tmp_path / "a.pptx"
    prs.save(f)
    secoes = ExtratorDeArquivos().extrair(f, "", "a.pptx")
    assert [(s.pagina, s.texto) for s in secoes] == [(1, "Slide um"), (2, "Slide dois")]


def test_html_and_text(tmp_path: Path) -> None:
    h = tmp_path / "a.html"
    h.write_text(
        "<html><title>T</title><body><p>Olá</p><script>x()</script></body></html>", encoding="utf-8"
    )
    t = tmp_path / "a.md"
    t.write_text("# Notas\nconteúdo")
    ex = ExtratorDeArquivos()
    assert ex.extrair(h, "text/html", "a.html")[0].texto == "Olá"
    assert ex.extrair(t, "", "a.md")[0].texto.startswith("# Notas")


def test_unsupported_format_raises(tmp_path: Path) -> None:
    f = tmp_path / "a.zip"
    f.write_bytes(b"PK")
    with pytest.raises(FormatoNaoSuportado):
        ExtratorDeArquivos().extrair(f, "application/zip", "a.zip")


def test_registry_is_case_insensitive() -> None:
    reg: Registry[int] = Registry()
    reg.register([".PDF"], 1)
    assert reg.get(".pdf") == 1 and reg.get(".doc") is None
